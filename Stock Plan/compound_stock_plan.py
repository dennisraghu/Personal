"""
Stock plan: 20% daily compound from $500 to $25,000.
Generates an Excel file with Date, Starting balance, Target for the day,
Planned closing balance, Actual Balance (to fill), Red or Green (formula).
"""

import math
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Configuration
STARTING_BALANCE = 500
TARGET_BALANCE = 25_000
DAILY_RATE = 0.20  # 20%


def trading_days_to_target(start: float, target: float, daily_pct: float) -> int:
    """Number of trading days to reach target with daily compound."""
    if start <= 0 or target <= start or daily_pct <= 0:
        return 0
    n = math.log(target / start) / math.log(1 + daily_pct)
    return math.ceil(n)


def next_trading_day(d: datetime) -> datetime:
    """Next weekday (no weekend)."""
    d = d + timedelta(days=1)
    while d.weekday() >= 5:  # 5=Saturday, 6=Sunday
        d += timedelta(days=1)
    return d


def main():
    n_days = trading_days_to_target(STARTING_BALANCE, TARGET_BALANCE, DAILY_RATE)
    print(f"Trading days to reach ${TARGET_BALANCE:,.0f}: {n_days}")

    # Build trading calendar (weekdays only)
    start_date = datetime.now().date()
    if start_date.weekday() >= 5:
        start_date = next_trading_day(datetime.combine(start_date, datetime.min.time())).date()

    balance = STARTING_BALANCE
    rows = []
    current_date = start_date

    for day in range(n_days):

        target_for_day = balance * DAILY_RATE
        planned_closing = balance * (1 + DAILY_RATE)

        date_str = current_date.strftime("%m/%d/%y")  # MM/DD/YY
        rows.append({
            "Date": date_str,
            "Starting balance": round(balance, 2),
            "Target for the day": round(target_for_day, 2),
            "Planned closing balance for the day": round(planned_closing, 2),
            "Actual Balance for the Day": None,
            "Red or Green": "",
        })
        balance = planned_closing
        current_date = next_trading_day(datetime.combine(current_date, datetime.min.time())).date()

    df = pd.DataFrame(rows)

    out_path = Path(__file__).parent / "stock_plan_20pct_daily.xlsx"
    df.to_excel(out_path, index=False, sheet_name="Stock Plan")

    # Add Red or Green formula (dates already written as MM/DD/YY strings)
    wb = load_workbook(out_path)
    ws = wb.active
    for row in range(2, len(rows) + 2):
        ws[f"F{row}"] = f'=IF(OR(ISBLANK(E{row}),E{row}=""),"",IF(E{row}>=D{row},"Green","Red"))'
    wb.save(out_path)

    print(f"Saved: {out_path}")
    print("Fill column 'Actual Balance for the Day' in Excel; 'Red or Green' will update automatically.")


if __name__ == "__main__":
    main()
